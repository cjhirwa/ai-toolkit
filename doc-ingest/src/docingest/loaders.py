from __future__ import annotations

import hashlib
import logging
import os
import re
from collections.abc import Iterator
from concurrent.futures import Future, ThreadPoolExecutor
from pathlib import Path

from .models import Document, PageSpan
from .ocr import OCRBackend

log = logging.getLogger("docingest")

TEXT_EXT = {".txt", ".md", ".markdown", ".rst"}
HTML_EXT = {".html", ".htm"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}
SUPPORTED = TEXT_EXT | HTML_EXT | IMAGE_EXT | {".pdf", ".docx"}


def _doc_id(data: bytes) -> str:
    return hashlib.sha1(data).hexdigest()[:16]


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # re-join hyphenated line breaks
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _from_pages(path: Path, data: bytes, pages: list[str], **meta) -> Document:
    text, spans = "", []
    for i, page_text in enumerate(pages, start=1):
        page_text = _clean(page_text)
        if not page_text:
            continue
        if text:
            text += "\n\n"
        spans.append(PageSpan(i, len(text), len(text) + len(page_text)))
        text += page_text
    return Document(_doc_id(data), str(path), text, spans, {"source": str(path), "title": path.stem, "n_pages": len(pages), **meta})


_CONTINUATION_WORDS = {"a", "an", "the", "and", "or", "but", "of", "to", "in", "on", "for", "with",
                       "is", "are", "was", "that", "which", "who", "as", "by", "at", "from"}


def _pdf_headings(pdf) -> list[dict[str, int]]:
    """Per page, {line text: heading level} for lines that look like headings.

    PDFs carry no heading markup, so this infers it from typography: the most common font
    size is body text; short lines set at least 15% larger are headings, ranked by size
    into levels 1-3, and a short standalone line that is entirely bold is the next level
    down. Ignored: sizes used on only one page (cover and title pages, which would
    otherwise become the parent of every heading), runs of more than 4 large lines
    (pull quotes), text repeated on many pages (running headers) and bare numbers.
    """
    lines = []  # (page index, block id, text, size, bold, alone in block)
    sizes: dict[float, int] = {}
    for i, page in enumerate(pdf):
        for b, block in enumerate(page.get_text("dict")["blocks"]):
            block_lines = block.get("lines", [])
            for line in block_lines:
                spans = [s for s in line["spans"] if s["text"].strip()]
                if not spans:
                    continue
                text = " ".join("".join(s["text"] for s in spans).split())
                size = round(max(s["size"] for s in spans) * 2) / 2
                sizes[size] = sizes.get(size, 0) + len(text)
                bold = all(s["flags"] & 16 for s in spans)
                lines.append((i, (i, b), text, size, bold, len(block_lines) == 1))
    if not sizes:
        return [{} for _ in range(len(pdf))]
    body = max(sizes, key=lambda s: sizes[s])

    def looks_like_heading(text: str) -> bool:
        last = text.rsplit(None, 1)[-1].lower() if text.strip() else ""
        return (2 <= len(text) <= 100 and any(c.isalpha() for c in text)
                and not text.endswith((".", ",", ";", "—", "-"))
                and last not in _CONTINUATION_WORDS)  # "… is that" is half a sentence, not a title

    pages_with: dict[str, set[int]] = {}
    for i, _, text, *_ in lines:
        pages_with.setdefault(text, set()).add(i)
    repeated = {t for t, p in pages_with.items() if len(p) >= max(3, len(pdf) // 10)}

    pages_of_size: dict[float, set[int]] = {}
    for i, _, text, size, _, _ in lines:
        if size >= body * 1.15 and text not in repeated and looks_like_heading(text):
            pages_of_size.setdefault(size, set()).add(i)
    big = sorted((s for s, p in pages_of_size.items() if len(p) >= 2), reverse=True)
    level_of = {s: min(rank + 1, 3) for rank, s in enumerate(big)}
    bold_level = min(len(big) + 1, 4)

    def runs(keep):
        """Consecutive lines on one page, at one size, that all satisfy ``keep``."""
        run: list = []
        for entry in lines + [None]:
            if entry is not None and keep(entry) and (not run or (run[-1][0], run[-1][3]) == (entry[0], entry[3])):
                run.append(entry)
                continue
            if run:
                yield run
            run = [entry] if entry is not None and keep(entry) else []

    out: list[dict[str, int]] = [{} for _ in range(len(pdf))]

    def mark(run, level):
        for e in run:
            out[e[0]].setdefault(e[2], level)

    # Large lines: up to 4 consecutive lines that together read as a title (so one wrapped
    # over lines ending in commas, "TO DO: CONSTRAINTS,", still counts); more is a pull quote.
    for run in runs(lambda e: e[3] in level_of and e[2] not in repeated):
        if len(run) <= 4 and looks_like_heading(" ".join(e[2] for e in run)):
            mark(run, level_of[run[0][3]])

    # Bold body-size lines that stand alone: one or two lines that do not read as a
    # sentence; longer runs are emphasised text.
    for run in runs(lambda e: e[4] and e[5] and body <= e[3] < body * 1.15 and e[2] not in repeated):
        joined = " ".join(e[2] for e in run)
        if len(run) <= 2 and len(joined) <= 80 and looks_like_heading(joined) and ". " not in joined:
            mark(run, bold_level)
    return out


def _mark_headings(page_text: str, headings: dict[str, int]) -> str:
    """Prefix heading lines with Markdown #s, joining a heading wrapped over several lines."""
    if not headings:
        return page_text
    out: list[str] = []
    prev_level = None
    for raw in page_text.split("\n"):
        text = " ".join(raw.split())
        level = headings.get(text)
        if level is None:
            out.append(raw)
            prev_level = None
        elif prev_level == level:
            out[-1] += " " + text  # second line of the same heading
        else:
            out.append(f"{'#' * level} {text}")
            prev_level = level
    return "\n".join(out)


def _check_page_count(path: Path, pages: int) -> None:
    limit = int(os.getenv("DOCINGEST_MAX_PAGES", "5000"))
    if 0 < limit < pages:
        raise ValueError(f"{path.name}: {pages} pages is over DOCINGEST_MAX_PAGES={limit}")


def load_pdf(path: Path, ocr: OCRBackend | None = None, min_chars: int = 25, headings: bool = True) -> Document:
    data = path.read_bytes()
    pages: list[str] = []
    ocr_pages: list[int] = []  # page numbers whose text came from OCR
    try:
        import pymupdf

        workers = getattr(ocr, "concurrency", 1)
        pending: dict[int, Future] = {}  # page number -> OCR result

        def collect(n: int) -> None:
            # One failed page (provider overloaded, timeout) must not cost the whole document.
            try:
                pages[n - 1] = pending.pop(n).result()
                ocr_pages.append(n)
            except Exception as e:
                log.warning("%s: OCR failed on page %d, leaving it out: %s", path.name, n, e)

        with pymupdf.open(stream=data, filetype="pdf") as pdf, ThreadPoolExecutor(max_workers=workers) as pool:
            _check_page_count(path, len(pdf))
            heads = _pdf_headings(pdf) if headings else [{} for _ in range(len(pdf))]
            for n, page in enumerate(pdf, start=1):
                t = _mark_headings(page.get_text("text"), heads[n - 1])
                pages.append(t)
                if len(t.strip()) < min_chars and ocr is not None:
                    # Render here (PyMuPDF is not thread-safe); only the model call runs in the pool.
                    pending[n] = pool.submit(ocr.image_to_text, page.get_pixmap(dpi=200).tobytes("png"))
                    if len(pending) >= 2 * workers:  # bound how many rendered pages wait in memory
                        collect(min(pending))
            for n in sorted(pending):
                collect(n)
        ocr_pages.sort()
    except ImportError:
        from pypdf import PdfReader

        # pypdf logs a warning per font it cannot fully decode (hundreds for one book)
        # even when the extracted text is fine; keep only its errors while we read.
        pypdf_log = logging.getLogger("pypdf")
        level = pypdf_log.level
        pypdf_log.setLevel(logging.ERROR)
        try:
            reader = PdfReader(path)
            _check_page_count(path, len(reader.pages))
            for page in reader.pages:
                pages.append(page.extract_text() or "")
        finally:
            pypdf_log.setLevel(level)
        if ocr is not None and any(len(p.strip()) < min_chars for p in pages):
            log.warning("%s: OCR needs PyMuPDF to render pages; install doc-ingest[pdf]. "
                        "Ingesting only the pages that have text.", path.name)
    empty = sum(1 for p in pages if len(p.strip()) < min_chars)
    if empty and ocr is None:
        log.warning("%s: %d/%d pages have no text layer; pass an OCR backend", path.name, empty, len(pages))
    return _from_pages(path, data, pages, ocr_pages=ocr_pages)


def load_docx(path: Path) -> Document:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    data = path.read_bytes()
    d = docx.Document(str(path))
    lines: list[str] = []
    body = d.element.body
    for el in body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, d)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style == "title":
                lines.append(f"# {text}")
            elif style.startswith("heading"):
                level = int(re.sub(r"\D", "", style) or 1)
                lines.append(f"{'#' * min(level, 6)} {text}")
            elif "list" in style:
                lines.append(f"- {text}")
            else:
                lines.append(text)
        elif tag == "tbl":
            t = Table(el, d)
            rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in t.rows]
            if rows:
                table = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
                table += ["| " + " | ".join(r) + " |" for r in rows[1:]]
                lines.append("\n".join(table))
    return Document(_doc_id(data), str(path), _clean("\n\n".join(lines)), [], {"source": str(path), "title": path.stem})


def load_html(path: Path) -> Document:
    from bs4 import BeautifulSoup

    data = path.read_bytes()
    soup = BeautifulSoup(data, "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else path.stem
    for tag in soup(["head", "title", "script", "style", "nav", "footer", "header", "noscript", "form"]):
        tag.decompose()
    for level in range(1, 7):
        for h in soup.find_all(f"h{level}"):
            h.replace_with(f"\n\n{'#' * level} {h.get_text(' ', strip=True)}\n\n")
    for li in soup.find_all("li"):
        li.replace_with(f"\n- {li.get_text(' ', strip=True)}")
    text = soup.get_text("\n")
    text = "\n".join(line.strip() for line in text.splitlines())
    return Document(_doc_id(data), str(path), _clean(text), [], {"source": str(path), "title": title})


def load_text(path: Path) -> Document:
    data = path.read_bytes()
    return Document(_doc_id(data), str(path), _clean(data.decode("utf-8", errors="replace")), [],
                    {"source": str(path), "title": path.stem})


def load_image(path: Path, ocr: OCRBackend | None) -> Document:
    if ocr is None:
        raise ValueError(f"{path.name} is an image; pass an OCR backend")
    data = path.read_bytes()
    mime = "image/jpeg" if path.suffix.lower() in (".jpg", ".jpeg") else f"image/{path.suffix.lower().lstrip('.')}"
    return _from_pages(path, data, [ocr.image_to_text(data, mime)], ocr_pages=[1])


def load(path: str | Path, ocr: OCRBackend | None = None) -> Document:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return load_pdf(path, ocr)
    if ext == ".docx":
        return load_docx(path)
    if ext in HTML_EXT:
        return load_html(path)
    if ext in TEXT_EXT:
        return load_text(path)
    if ext in IMAGE_EXT:
        return load_image(path, ocr)
    raise ValueError(f"Unsupported file type: {path}")


def iter_files(paths: list[str | Path], pattern: str = "**/*") -> Iterator[Path]:
    for p in map(Path, paths):
        if p.is_dir():
            yield from sorted(f for f in p.glob(pattern) if f.is_file() and f.suffix.lower() in SUPPORTED)
        elif p.exists():
            yield p
        else:
            raise FileNotFoundError(p)
